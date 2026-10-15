
from pathlib import Path

from openpyxl import load_workbook

from app.tools.spreadsheet import EXCEL_FILE
from app.tools.email_drafts import DRAFTS_DIR


def verify_results(context: dict) -> dict:
    """Verify internship results and email drafts for the current task."""

    internships = context.get("internships", [])
    drafts = context.get("drafts", [])

    spreadsheet_exists = EXCEL_FILE.exists()
    spreadsheet_rows = 0
    spreadsheet_error = None

    if spreadsheet_exists:
        workbook = None
        try:
            workbook = load_workbook(EXCEL_FILE, read_only=True)

            if "Internships" not in workbook.sheetnames:
                spreadsheet_error = "Internships worksheet is missing."
            else:
                sheet = workbook["Internships"]
                spreadsheet_rows = max(sheet.max_row - 1, 0)

        except Exception as exc:
            spreadsheet_error = str(exc)

        finally:
            if workbook is not None:
                workbook.close()

    # Check that each current search result has a matching spreadsheet row.
    saved_keys = set()
    if spreadsheet_exists and spreadsheet_error is None:
        workbook = load_workbook(EXCEL_FILE, read_only=True)
        try:
            sheet = workbook["Internships"]

            headers = [
                str(cell.value or "").strip()
                for cell in sheet[1]
            ]
            company_col = headers.index("Company") + 1 if "Company" in headers else None
            role_col = headers.index("Role") + 1 if "Role" in headers else None
            id_col = (
                headers.index("Source Listing ID") + 1
                if "Source Listing ID" in headers
                else None
            )

            for row in sheet.iter_rows(min_row=2, values_only=True):
                company = row[company_col - 1] if company_col and len(row) >= company_col else ""
                role = row[role_col - 1] if role_col and len(row) >= role_col else ""
                listing_id = row[id_col - 1] if id_col and len(row) >= id_col else ""

                if listing_id:
                    saved_keys.add(("id", str(listing_id).strip()))
                if company and role:
                    saved_keys.add((
                        "company_role",
                        str(company).strip().lower(),
                        str(role).strip().lower(),
                    ))
        finally:
            workbook.close()

    missing_internships = []
    for item in internships:
        listing_id = item.get("source_listing_id")
        company = str(item.get("company") or "").strip().lower()
        role = str(item.get("role") or "").strip().lower()

        found = (
            ("id", str(listing_id).strip()) in saved_keys
            if listing_id
            else False
        )

        if not found and company and role:
            found = ("company_role", company, role) in saved_keys

        if not found:
            missing_internships.append(
                item.get("role") or item.get("company") or "Unknown internship"
            )

    missing_drafts = []
    for draft in drafts:
        file_path = draft.get("file")
        if not file_path or not Path(file_path).is_file():
            missing_drafts.append(file_path or "Draft has no file path")

    # Do not claim the workflow succeeded if no internships were found.
    verified = (
        spreadsheet_exists
        and spreadsheet_error is None
        and len(internships) > 0
        and not missing_internships
        and not missing_drafts
    )

    return {
        "verified": verified,
        "spreadsheet_exists": spreadsheet_exists,
        "spreadsheet_rows": spreadsheet_rows,
        "internships_found": len(internships),
        "internships_missing_from_sheet": missing_internships,
        "drafts_created": len(drafts),
        "missing_drafts": missing_drafts,
        "spreadsheet_error": spreadsheet_error,
    }


if __name__ == "__main__":
    print(verify_results({"internships": [], "drafts": []}))
